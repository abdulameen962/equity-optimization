"""Catalogue all formula/formatting issues in the DOCX."""
import sys
sys.stdout.reconfigure(encoding='utf-8')

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = Document('Abdulameen_Complete_Thesis_Chapters_1_5.docx')

# Search for key phrases to locate each issue
searches = [
    ("Feature vector", "feature vector.*13.*1"),
    ("Node split Rnode", "node.*region.*feature space"),
    ("Rockafellar CVaR", "Rockafellar and Uryasev.*2000.*CVaR"),
    ("CVaR optimization objective", "minimize the tail risk.*CVaR.*minimum"),
    ("Turnover formula", "Turnover"),
    ("Transaction costs", "Following standard empirical finance"),
    ("MAE formula", "Mean Absolute Error"),
    ("Sharpe ratio", "Sharpe"),
    ("Sortino ratio", "Sortino"),
    ("Maximum Drawdown", "Maximum Drawdown"),
    ("Buy-and-Hold weights", "Buy-and-Hold.*weights drift"),
    ("Hypothesis Testing 3.3.6", "Hypothesis Testing"),
    ("Variable table", "tabular summary of all variables"),
    ("Table 4.3", "Table 4.3"),
    ("References", "References"),
    ("Risk-free rate Rf", "risk-free rate proxy"),
]

import re

for label, pattern in searches:
    print(f"\n=== {label} ===")
    for i, p in enumerate(doc.paragraphs):
        txt = p.text.strip()
        if re.search(pattern, txt, re.IGNORECASE):
            align = p.alignment
            safe = txt[:120].encode('ascii', 'replace').decode('ascii')
            print(f"  P{i} [{align}]: {safe}")
            # Show surrounding paragraphs
            for j in range(max(0,i-1), min(len(doc.paragraphs), i+4)):
                stxt = doc.paragraphs[j].text.strip()[:100].encode('ascii', 'replace').decode('ascii')
                marker = ">>>" if j == i else "   "
                print(f"  {marker} P{j}: {stxt}")
            break

# Check for repeated headers
print("\n=== REPEATED HEADERS ===")
header_texts = {}
for i, p in enumerate(doc.paragraphs):
    txt = p.text.strip()
    if txt and len(txt) < 80:
        is_heading = (
            txt.startswith('Chapter') or txt.startswith('CHAPTER') or
            re.match(r'^\d+\.\d+', txt) or
            txt == 'LITERATURE REVIEW' or txt == 'RESEARCH METHODOLOGY' or
            txt == 'References'
        )
        if is_heading:
            if txt in header_texts:
                print(f"  DUPLICATE: '{txt}' at P{header_texts[txt]} and P{i}")
            header_texts[txt] = i

# Check where References section starts
print("\n=== REFERENCES LOCATION ===")
for i, p in enumerate(doc.paragraphs):
    if p.text.strip() == 'References':
        print(f"  P{i}: References heading")
        # Check paragraph before it
        if i > 0:
            prev = doc.paragraphs[i-1].text.strip()[:80]
            print(f"  P{i-1} (before): {prev}")
